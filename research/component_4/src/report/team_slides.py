"""Component 4's three slides for the team PP1 deck (J26-IT-363), in that deck's style: the module title slide and
the two required content slides (sub-problem / prototype, design excellence + user feedback).
    .venv\\Scripts\\python.exe -m src.report.team_slides
Writes docs/slides/PP1_team_component4.pptx. Numbers are read from results/ (same sources as src.report.slides).
Copy the three slides into the team deck (Home -> Reuse slides, or copy / paste with "Keep source formatting").
"""
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from src.report.slides import FIG, OUT, RES, RES_AP, bench, latest

# team deck palette (read off J26-IT-363_PP1.pdf)
DARK, CARD_DARK, GREY_BG, INK, MUTED = "0F332B", "15382F", "F2F2F2", "1F2A24", "6B6B6B"
GOLD, BLUE, GREEN, RED = "D4A12A", "2A7FA3", "4C8A55", "C95C44"
TINT = {"blue": "E2EDF4", "green": "E5F0E3", "yellow": "F6EBD0", "red": "F4E0DC"}
CREAM, CREAM_LINE = "FFF7E6", "E8C77A"
SERIF, SANS = "Georgia", "Segoe UI"
W, H = Inches(13.333), Inches(7.5)
FOOTER = "IT23343948 | Weerasekera J.V | J26-IT-363"


def rgb(h):
    return RGBColor.from_string(h)


def box(s, x, y, w, h, fill, line=None, radius=0.06):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    b.adjustments[0] = radius
    b.fill.solid()
    b.fill.fore_color.rgb = rgb(fill)
    if line:
        b.line.color.rgb = rgb(line)
        b.line.width = Pt(1)
    else:
        b.line.fill.background()
    b.shadow.inherit = False
    return b


def text(s, x, y, w, h, runs, size=14, color=INK, font=SANS, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=4):
    """runs: str (lines split on \\n) or list of lines; a line may be [(text, {bold, color, size, font}), ...]."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.02)
    lines = runs.split("\n") if isinstance(runs, str) else runs
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        for t, o in ([(line, {})] if isinstance(line, str) else line):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name, f.size, f.bold = o.get("font", font), Pt(o.get("size", size)), o.get("bold", bold)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def label(s, x, y, t, color, w=Inches(6)):
    return text(s, x, y, w, Inches(0.3), t.upper(), size=10, color=color, bold=True)


def page_number(s, n, color=MUTED):
    text(s, W - Inches(0.6), H - Inches(0.4), Inches(0.4), Inches(0.3), str(n), size=8, color=color, align=PP_ALIGN.RIGHT)


def phone(s, path, x, y, h):
    pic = s.shapes.add_picture(str(path), x, y, height=h)
    pic.line.color.rgb = rgb("C9D3CD")
    pic.line.width = Pt(1)
    return pic


def numbers() -> dict:
    hp_t, hp_s = bench(RES)
    q = {k: pd.read_csv(latest(f, "LATEST_QUOTA") / "summary.csv").set_index("rule") for k, f in (("hp", RES), ("ap", RES_AP))}
    cut = {k: 100 * (1 - v.loc["stgnn", "person_hours_over"] / v.loc["static", "person_hours_over"]) for k, v in q.items()}
    ctl = pd.read_csv(latest(RES, "LATEST_CONTROL") / "summary.csv")
    slots = 100 * (1 - ctl.person_hours_over_capacity[3] / ctl.person_hours_over_capacity[0])
    sp = {m: hp_s.loc[m, ("MAE", "mean")] for m in ("HA", "lstm", "gru_nograph", "stgnn")}
    tp = {m: hp_t.loc[m, "MAE"] for m in ("HA", "ARIMA", "lstm", "gru_nograph", "stgnn")}
    return {"cut_hp": cut["hp"], "cut_ap": cut["ap"], "slots": slots, "sp": sp, "tp": tp}


def build():
    n = numbers()
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    # ------------------------------------------------------------------ 15 · module title
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(DARK)
    bg.line.fill.background()
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.62), Inches(1.0), Inches(0.72), Inches(0.72))
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(BLUE)
    c.line.fill.background()
    text(s, Inches(0.62), Inches(1.0), Inches(0.72), Inches(0.72), "4", size=22, color="FFFFFF", bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    label(s, Inches(1.52), Inches(1.02), "Module 4", BLUE)
    text(s, Inches(1.52), Inches(1.3), Inches(8.65), Inches(1.2),
         "Spatio-Temporal Bottleneck Forecasting &\nDynamic Carrying Capacity", size=26, color="FFFFFF", font=SERIF, bold=True,
         spacing=0)
    text(s, Inches(1.52), Inches(2.55), Inches(8.4), Inches(0.8),
         "Forecasts where and when trail sections overload, a day ahead, from open data only - and turns it into a "
         "safe number per trail instead of one fixed daily quota.", size=14, color="D8E3DD")
    ph = box(s, Inches(10.27), Inches(0.85), Inches(2.06), Inches(2.45), "1E4A40", radius=0.02)
    ph.text_frame.text = "Your photo\n(keep the one in the team deck)"
    for p in ph.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size, r.font.color.rgb, r.font.name = Pt(11), rgb("9FB8AE"), SANS
    box(s, Inches(0.62), Inches(3.4), Inches(11.7), Inches(2.2), CARD_DARK, radius=0.03)
    label(s, Inches(0.92), Inches(3.68), "Contribution to the overall solution", BLUE)
    text(s, Inches(0.92), Inches(4.05), Inches(11.1), Inches(0.9),
         "Supplies the crowd and capacity signal of the framework: a per-section forecast and safe numbers that the "
         "safety assistant (Module 3) and the authority dashboard read offline through a versioned data contract.",
         size=15, color="FFFFFF")
    text(s, Inches(0.92), Inches(5.0), Inches(11), Inches(0.35),
         [[("SPECIALIZATION      ", {"size": 10, "bold": True, "color": BLUE}),
           ("Information Technology", {"size": 13, "bold": True, "color": "FFFFFF"})]])
    text(s, Inches(0.62), Inches(6.9), Inches(6), Inches(0.3), FOOTER, size=9, color="9FB8AE")

    # ------------------------------------------------------------------ 16 · sub-problem / prototype
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(GREY_BG)
    bg.line.fill.background()
    label(s, Inches(0.55), Inches(0.62), "Module 4 · Crowd & capacity forecasting", GOLD)
    text(s, Inches(0.55), Inches(0.9), Inches(12.2), Inches(0.7), "How the solution addresses the sub-problem / prototype",
         size=28, font=SERIF, bold=True)
    text(s, Inches(0.55), Inches(1.55), Inches(12), Inches(0.35),
         "Working prototype at three parks (Horton Plains, Adam's Peak, Yala). Results: 2025 replay, simulation.",
         size=12, color=MUTED)

    # left: sub-problem + requirements
    box(s, Inches(0.55), Inches(2.0), Inches(3.75), Inches(4.75), "FFFFFF")
    label(s, Inches(0.8), Inches(2.2), "Sub-problem", BLUE)
    text(s, Inches(0.8), Inches(2.5), Inches(3.3), Inches(1.4),
         "One fixed daily quota per park cannot say when or where a trail section overloads - and no section counts "
         "exist.", size=15, font=SERIF, bold=True)
    label(s, Inches(0.8), Inches(3.95), "User requirements met", BLUE)
    text(s, Inches(0.8), Inches(4.25), Inches(3.3), Inches(2.4), [
        [("✓  ", {"color": GREEN, "bold": True}), ("Section-level forecast, next 24 h (FR2)", {})],
        [("✓  ", {"color": GREEN, "bold": True}), ("Condition-aware capacity, Cifuentes (FR3)", {})],
        [("✓  ", {"color": GREEN, "bold": True}), ("Runs offline on the phone < 2 s (NFR1)", {})],
        [("✓  ", {"color": GREEN, "bold": True}), ("One number + 'why' for officials (NFR4)", {})],
        [("✓  ", {"color": GREEN, "bold": True}), ("Feed for Module 3 + dashboard (FR4)", {})],
        [("✓  ", {"color": GREEN, "bold": True}), ("Tested against simple baselines", {})],
    ], size=11.5, spacing=5)

    # middle: the real pipeline
    box(s, Inches(4.5), Inches(2.0), Inches(3.55), Inches(4.75), TINT["blue"])
    label(s, Inches(4.75), Inches(2.2), "Prototype flow", BLUE)
    steps = [("OpenStreetMap trail graph", "35-57 sections of ~300 m - 1.2 km"),
             ("Terrain + weather + calendar", "Copernicus DEM, Open-Meteo, Poya days"),
             ("ST-GNN: GCN + GRU", "people per section per hour, on the phone"),
             ("Cifuentes PCC → RCC → ECC", "forecast enters the real-capacity step"),
             ("Decisions", "safe numbers · entry slots · gate advice")]
    for i, (a, b) in enumerate(steps):
        y = Inches(2.55) + i * Inches(0.82)
        box(s, Inches(4.75), y, Inches(3.05), Inches(0.68), "FFFFFF", radius=0.2)
        text(s, Inches(4.85), y + Inches(0.05), Inches(2.85), Inches(0.62),
             [[(a, {"bold": True, "size": 11.5})], [(b, {"size": 9.5, "color": MUTED})]], spacing=0,
             align=PP_ALIGN.CENTER)

    # right: prototype screens + results
    shots = OUT / "screens"
    phone(s, shots / "11_bottleneck_map.png", Inches(8.3), Inches(2.0), Inches(3.35))
    phone(s, FIG / "15_daily_quota_cropped.png" if (FIG / "15_daily_quota_cropped.png").exists()
          else shots / "15_daily_quota.png", Inches(9.95), Inches(2.0), Inches(3.35))
    text(s, Inches(8.3), Inches(5.4), Inches(1.6), Inches(0.5), "Tomorrow, section by section (offline)",
         size=9, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, Inches(9.95), Inches(5.4), Inches(1.6), Inches(0.5), "Officer: forecast vs fixed quota",
         size=9, color=MUTED, align=PP_ALIGN.CENTER)
    box(s, Inches(11.75), Inches(2.0), Inches(1.05), Inches(4.75), DARK, radius=0.1)
    for i, (v, t_) in enumerate([(f"−{n['cut_hp']:.0f}%", "over-capacity\nvs fixed quota\n(Horton)"),
                                 (f"−{n['cut_ap']:.0f}%", "same, at\nAdam's Peak"),
                                 (f"−{n['slots']:.1f}%", "with timed\nentry slots"),
                                 ("<2 s", "forecast on\nthe phone")]):
        y = Inches(2.15) + i * Inches(1.15)
        text(s, Inches(11.75), y, Inches(1.05), Inches(1.1),
             [[(v, {"size": 17, "bold": True, "color": GOLD, "font": SERIF})], [(t_, {"size": 8, "color": "D8E3DD"})]],
             align=PP_ALIGN.CENTER, spacing=0)
    text(s, Inches(8.3), Inches(5.95), Inches(3.4), Inches(0.8),
         "Same visitors admitted in every comparison. Slot tickets fix a trail + entry time, so no section goes over its "
         "safe number in any half hour.", size=9, color=MUTED)
    page_number(s, 16)

    # ------------------------------------------------------------------ 17 · design excellence + feedback
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(GREY_BG)
    bg.line.fill.background()
    label(s, Inches(0.55), Inches(0.62), "Module 4 · Crowd & capacity forecasting", GOLD)
    text(s, Inches(0.55), Inches(0.9), Inches(12.2), Inches(0.7), "Design excellence / contribution + user feedback",
         size=28, font=SERIF, bold=True)
    text(s, Inches(0.55), Inches(1.55), Inches(12), Inches(0.35),
         "Tested, not assumed: the graph model is compared with simpler models, and the parts that did not work are "
         "reported.", size=12, color=MUTED)

    cards = [
        ("Graph structure", BLUE, "blue",
         "Tiny trail graphs (35-57 sections) built from OpenStreetMap - not dense city traffic. One config file per "
         "park: Horton Plains, Adam's Peak, Yala."),
        ("Transfer strategy", GREEN, "green",
         "Pretrained on METR-LA → PEMS-BAY, adapted to each trail. Honest result: helps only with 30-90 days of data; "
         "traffic weights alone fail."),
        ("Management value", RED, "red",
         f"Forecast enters Cifuentes RCC → a safe number per trail that changes with weather and crowding: "
         f"−{n['cut_hp']:.0f}% / −{n['cut_ap']:.0f}% over-capacity time vs a fixed quota."),
    ]
    for i, (t_, col, tint, body) in enumerate(cards):
        x = Inches(0.55) + i * Inches(2.62)
        box(s, x, Inches(2.0), Inches(2.47), Inches(2.05), TINT[tint])
        text(s, x + Inches(0.18), Inches(2.15), Inches(2.15), Inches(0.3), t_, size=11, color=col, bold=True)
        text(s, x + Inches(0.18), Inches(2.5), Inches(2.15), Inches(1.5), body, size=10.5)

    # validation: baseline table + forecast plot
    box(s, Inches(0.55), Inches(4.2), Inches(7.71), Inches(2.6), TINT["yellow"])
    text(s, Inches(0.75), Inches(4.3), Inches(4), Inches(0.3), "Validation · baseline comparison", size=11, color="B8860B",
         bold=True)
    tp, sp = n["tp"], n["sp"]
    rows = [["MAE (people / section-hour)", "Hist. avg", "ARIMA", "LSTM", "GRU, no graph", "ST-GNN"],
            ["Every section counted", f"{tp['HA']:.2f}", f"{tp['ARIMA']:.2f}", f"{tp['lstm']:.2f}", f"{tp['gru_nograph']:.2f}",
             f"{tp['stgnn']:.2f}"],
            ["Sections without counts", f"{sp['HA']:.2f}", "–", f"{sp['lstm']:.2f}", f"{sp['gru_nograph']:.2f}",
             f"{sp['stgnn']:.2f}"]]
    tbl = s.shapes.add_table(3, 6, Inches(0.75), Inches(4.65), Inches(4.3), Inches(1.05)).table
    for j, wdt in enumerate([1.45, 0.55, 0.5, 0.5, 0.65, 0.65]):
        tbl.columns[j].width = Inches(wdt)
    for r, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = tbl.cell(r, j)
            cell.text = v
            cell.margin_left = cell.margin_right = Inches(0.04)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            f = p.runs[0].font
            f.size, f.name = Pt(8.5), SANS
            f.bold = r == 0 or (j == 5 and r == 2)
            f.color.rgb = rgb("FFFFFF" if r == 0 else INK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(BLUE if r == 0 else "FFFFFF")
    text(s, Inches(0.75), Inches(5.78), Inches(4.3), Inches(0.95),
         "Horton Plains, 2025 test year, 3 seeds, paired tests over days. Graph helps on uncounted sections "
         "(p < 0.001, ~2 %); a tie when every section is counted. Spatial block CV; no leakage.", size=8.5, color=MUTED)
    s.shapes.add_picture(str(FIG / "example_day.png"), Inches(5.15), Inches(4.45), width=Inches(3.0))
    text(s, Inches(5.15), Inches(5.78), Inches(3.0), Inches(0.9),
         "Day-ahead forecast vs actual at World's End: right on typical days, low on the rare busiest day → decisions "
         "use the busy case.", size=8.5, color=MUTED)

    # user feedback (to fill with real answers before presenting)
    box(s, Inches(8.45), Inches(2.0), Inches(4.35), Inches(3.55), CREAM, line=CREAM_LINE)
    text(s, Inches(8.7), Inches(2.15), Inches(3.9), Inches(0.35), "USER FEEDBACK", size=13, color="7A5C12", bold=True)
    text(s, Inches(8.7), Inches(2.55), Inches(3.9), Inches(2.9), [
        [("Fill with real answers before presenting:", {"size": 10, "color": MUTED})],
        [("• Hiker: ", {"bold": True}), ("would you change your visit if a section shows red at 08:00?", {})],
        [("• Guide / tour operator: ", {"bold": True}), ("is the safe number + reason clear?", {})],
        [("• Ranger / DWC contact: ", {"bold": True}), ("would the counting + gate advice fit the job?", {})],
        [("• Changed because of it: ", {"bold": True}), ("…", {})],
    ], size=10, spacing=6)
    box(s, Inches(8.45), Inches(5.7), Inches(4.35), Inches(1.1), DARK, radius=0.1)
    text(s, Inches(8.65), Inches(5.8), Inches(4.0), Inches(0.95), [
        [("Next  ", {"bold": True, "color": GOLD}),
         ("Field counts in December (criteria fixed in advance) to validate the simulated target; surveys with DWC, "
          "rangers and hikers.", {"color": "FFFFFF"})]], size=10)
    page_number(s, 17)

    path = OUT / "PP1_team_component4.pptx"
    prs.save(path)
    return path


if __name__ == "__main__":
    print(build())
